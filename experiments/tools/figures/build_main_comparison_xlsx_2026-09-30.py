"""Build one Excel workbook from the main-comparison tables (2026-09-30, user request; English version, 18h added).
Inputs: experiments/analysis/figures/main_comparison_*_best10seeds*_2026-09-30_en.json (best checkpoint x evaluation seeds; same source as the PNGs)
        + Dublin 11h tail-40 block (constants S/EV in plot_main_table_en_2026-09-22.py = the 11h block of main_comparison_8std_vs_frap_2026-09-22_en.png).
Output: experiments/analysis/tables/main_comparison_tables_2026-09-30.xlsx
Means / s.d. are measured data (from the json); relative differences, ordering gates, type-swap statistics and per-seed re-computations are formulas. Font Arial."""
import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
REPO = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
FIG = os.path.join(REPO, "experiments/analysis/figures"); OUT = os.path.join(REPO, "experiments/analysis/tables/main_comparison_tables_2026-09-30.xlsx")
F = "Arial"; B = Font(name=F, bold=True); N = Font(name=F); TITLE = Font(name=F, bold=True, size=12); NOTE = Font(name=F, italic=True, color="555555")
HFILL = PatternFill("solid", fgColor="C9D3DF"); TFILL = PatternFill("solid", fgColor="E3E8EF"); GREEN = PatternFill("solid", fgColor="CFE9CF")
thin = Side(style="thin", color="999999"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
HDR = ["Metric", "8STD mean", "8STD s.d.", "GS-ENUM mean", "GS-ENUM s.d.", "Rel. diff (GS-8STD)/8STD", "GS wins / paired", "Notes"]

def style_row(ws, r, ncol=8, bold=False, fill=None):
    for c in range(1, ncol + 1):
        cell = ws.cell(r, c); cell.font = B if bold else N; cell.border = BOX; cell.alignment = Alignment(vertical="center", wrap_text=(c in (1, 8)))
        if fill: cell.fill = fill

def write_block(ws, r, title, rows_t, rows_e, amb_na, note=""):
    """One block: title + header + 5 stopped-time rows + gate + 5 stop-event rows + gate. Returns (next free row, key cells)."""
    ws.cell(r, 1, title).font = TITLE; ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8); ws.cell(r, 1).fill = TFILL; r += 1
    for c, h in enumerate(HDR, 1): ws.cell(r, c, h)
    style_row(ws, r, bold=True, fill=HFILL); r += 1
    key = {}
    def rows(block, tag):
        nonlocal r
        for row in block:
            lab, am, asd, bm, bsd, _rel, wins = (row + [""] * 7)[:7]
            ws.cell(r, 1, lab); ws.cell(r, 2, float(am)); ws.cell(r, 3, float(asd)); ws.cell(r, 4, float(bm)); ws.cell(r, 5, float(bsd))
            ws.cell(r, 6, f'=IF(B{r}=0,"",D{r}/B{r}-1)'); ws.cell(r, 7, wins if "J(531)" in lab else "")
            for c in (2, 3, 4, 5): ws.cell(r, c).number_format = "0.00"
            ws.cell(r, 6).number_format = "+0%;-0%;0%"
            style_row(ws, r)
            ws.cell(r, 2 if float(am) <= float(bm) else 4).fill = GREEN   # lower is better
            if "ambulance" in lab: ws.cell(r, 8, "Few ambulances per evaluation in Dublin (02h: 2, 11h: 4, 18h: 1); do not cite the mean alone")
            key[(tag, lab.split(" - ")[-1] if " - " in lab else "J")] = r
            r += 1
    rows(rows_t, "t")
    car, bus, amb = key[("t", "car")], key[("t", "bus")], key[("t", "ambulance")]
    ws.cell(r, 1, "Ordering gate (stopped time): amb <= bus <= car")
    for c, col in ((2, "B"), (4, "D")):
        ws.cell(r, c, f'=IF(AND({col}{amb}<={col}{bus},{col}{bus}<={col}{car}),"pass",IF(AND({col}{amb}-{col}{bus}<=5,{col}{bus}-{col}{car}<=5),"pass (tol. 5 s)","fail"))')
    ws.cell(r, 8, "tol. 5 s = tolerance of one 5 s sampling quantum"); style_row(ws, r); r += 1
    rows(rows_e, "e")
    car, bus, amb = key[("e", "car")], key[("e", "bus")], key[("e", "ambulance")]
    ws.cell(r, 1, "Ordering gate (stop events): amb <= bus <= car" + (" (amb n/a: bus <= car only)" if amb_na else ""))
    for c, col in ((2, "B"), (4, "D")):
        if amb_na: ws.cell(r, c, f'=IF({col}{bus}<={col}{car},"pass (amb n/a)","fail (amb n/a)")')
        else: ws.cell(r, c, f'=IF(AND({col}{amb}<={col}{bus},{col}{bus}<={col}{car}),"pass",IF(AND({col}{amb}-{col}{bus}<=0.5,{col}{bus}-{col}{car}<=0.5),"pass (tol.)","fail"))')
    if amb_na: ws.cell(r, 8, "1-4 ambulances per evaluation: stop events take values 0 / 0.5 / 1 only, not comparable with the bus mean over hundreds of vehicles")
    style_row(ws, r); r += 1
    if note: ws.cell(r, 1, note).font = NOTE; ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8); r += 1
    return r + 1, key

def write_typeswap(ws, r, seeds, ts, title):
    """Same route & departure time, ambulance vs bus probe: per-seed raw values + formula mean / s.d. / counts."""
    ws.cell(r, 1, title).font = TITLE; ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8); ws.cell(r, 1).fill = TFILL; r += 1
    hdr = ["Vehicle / metric"] + [f"seed {s}" for s in seeds] + ["Mean", "S.d."]
    for c, h in enumerate(hdr, 1): ws.cell(r, c, h)
    style_row(ws, r, ncol=len(hdr), bold=True, fill=HFILL); r += 1
    n = len(seeds); last = get_column_letter(1 + n); lines = []
    for arm in ("8STD", "GS"):
        for metric, mk in (("Stopped time / visit (s)", "pv"), ("Stop events / visit", "ev")):
            for veh, idx in (("ambulance (probe route)", 0), ("bus (same route & time)", 1)):
                ws.cell(r, 1, f"{arm} - {veh} - {metric}")
                for k, v in enumerate(ts[arm][mk][idx]): ws.cell(r, 2 + k, float(v)).number_format = "0.00"
                ws.cell(r, 2 + n, f"=AVERAGE(B{r}:{last}{r})").number_format = "0.00"; ws.cell(r, 3 + n, f"=STDEV(B{r}:{last}{r})").number_format = "0.00"
                style_row(ws, r, ncol=len(hdr)); lines.append((arm, mk, idx, r)); r += 1
    for arm in ("8STD", "GS"):
        for metric, mk in (("stopped time", "pv"), ("stop events", "ev")):
            ra = [x[3] for x in lines if x[0] == arm and x[1] == mk and x[2] == 0][0]; rb = [x[3] for x in lines if x[0] == arm and x[1] == mk and x[2] == 1][0]
            ws.cell(r, 1, f"{arm} - {metric}: evaluations with ambulance < bus / ties")
            ws.cell(r, 2, f'=SUMPRODUCT(--(B{ra}:{last}{ra}<B{rb}:{last}{rb}))&" / "&SUMPRODUCT(--(B{ra}:{last}{ra}=B{rb}:{last}{rb}))&" (of {n})"')
            style_row(ws, r, ncol=2); r += 1
    return r + 1

def block_pairs(d):
    out = {}
    for k, v in d.items():
        if "|" in k and isinstance(v, list):
            t, m = k.split("|"); out.setdefault(t, {})[m] = v
    return [(t, b.get("avg_stopped_time_per_visit"), b.get("avg_stop_events_per_visit")) for t, b in out.items()]

def add_json(wb, sheet, path, amb_na, protocol, title_override=None, ws=None, r=None, typeswap=True, ts_title=None, per_seed=True):
    """Append the blocks of one json to a sheet (created if ws is None). title_override may be a str (all blocks) or a list (one per block)."""
    d = json.load(open(path))
    if ws is None: ws = wb.create_sheet(sheet); r = 1
    ws.cell(r, 1, f"{sheet}: best checkpoints x {len(d['seeds'])} evaluation seeds ({', '.join(map(str, d['seeds']))}); {protocol}").font = TITLE; r += 1
    ws.cell(r, 1, f"Source: {os.path.relpath(path, REPO)} (same data as the PNG of that name); anomalous seeds (symmetric rule): {d.get('anomalous') or 'none'}; Tukey outlier seeds: {d.get('tukey_outliers') or 'none'}").font = NOTE; r += 2
    keys = []
    for i, (t, rt, re_) in enumerate(block_pairs(d)):
        tt = (title_override[i] if isinstance(title_override, list) else title_override) or t
        r, key = write_block(ws, r, tt, rt, re_, amb_na); keys.append((tt, key))
    if typeswap and "typeswap" in d:
        r = write_typeswap(ws, r, d["seeds"], d["typeswap"], ts_title or f"{sheet}: same route & departure time, ambulance vs bus probe, best checkpoints, {len(d['seeds'])} seeds")
    if per_seed and "per_seed" in d:
        ws.cell(r, 1, f"{sheet}: per-seed raw values, mean / s.d. recomputed by formula").font = TITLE; ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8); ws.cell(r, 1).fill = TFILL; r += 1
        seeds = d["seeds"]; n = len(seeds); last = get_column_letter(1 + n)
        hdr = ["Arm / class / metric"] + [f"seed {s}" for s in seeds] + ["Mean", "S.d."]
        for c, h in enumerate(hdr, 1): ws.cell(r, c, h)
        style_row(ws, r, ncol=len(hdr), bold=True, fill=HFILL); r += 1
        for arm in ("8STD", "GS"):
            for cls in ("car", "bus", "ambulance", "all"):
                for m, ml in (("avg_stopped_time_per_visit", "stopped time / visit (s)"), ("avg_stop_events_per_visit", "stop events / visit")):
                    vals = d["per_seed"][arm][cls][m]; ws.cell(r, 1, f"{arm} - {cls} - {ml}")
                    for k, v in enumerate(vals): ws.cell(r, 2 + k, float(v)).number_format = "0.00"
                    ws.cell(r, 2 + n, f"=AVERAGE(B{r}:{last}{r})").number_format = "0.00"; ws.cell(r, 3 + n, f"=STDEV(B{r}:{last}{r})").number_format = "0.00"
                    style_row(ws, r, ncol=len(hdr)); r += 1
        r += 1
    ws.column_dimensions["A"].width = 50
    for c in "BCDEFG": ws.column_dimensions[c].width = 16
    ws.column_dimensions["H"].width = 62; ws.freeze_panes = "B4"
    return ws, r + 1, keys


import glob, numpy as np
PV = os.path.join(REPO, "experiments/analysis/data/pervehicle")
SH18 = {"car": .9568, "bus": .0430, "ambulance": .0002}; WT = {"car": 1, "bus": 3, "ambulance": 5}

def _load(pat):
    out = {}
    for f in glob.glob(pat):
        seed = int(f.split("_seed")[1].split("_")[0]); out[seed] = json.load(open(f))
    return out

def _pv(recs, field):
    vals = [sum(r[field].values()) / r["k_v"] for r in recs if r["k_v"] > 0]; return float(np.mean(vals)) if vals else float("nan")

def three_arm_block(ws, r, title, arms, thr=100.0, seeds_keep=None, amb_ids=("amb_1",), bus_ids=("probe_bus_1",)):
    """arms = [(label, exp, ep), ...] (first = 8STD reference). Values computed from the raw collector json exactly like plot_best10seeds_table
    (population s.d.; J = sum SH*WT*metric; wins = seeds where the GS arm is lower than 8STD). Returns (next row, key cells, anomalous seeds)."""
    A = [_load(f"{PV}/exp{e}_ep{ep:05d}_seed*_best_coll.json") for _, e, ep in arms]
    Ab = [_load(f"{PV}/exp{e}_ep{ep:05d}_seed*_best_busprobe_coll.json") for _, e, ep in arms]
    seeds = sorted(set.intersection(*[set(x) for x in A + Ab]))
    bad = sorted({s for s in seeds for X in A if any(v > thr for k, v in X[s]["collector"]["summary"].items() if k.endswith("/all/avg_stopped_time") and not k.startswith("eval/system"))})
    if seeds_keep is not None: seeds = [s for s in seeds if s in seeds_keep]
    n = len(arms); ncol = 1 + 2 * n + (n - 1) + (1 if n > 2 else 0) + (n - 1) + 1
    hdr = ["Metric"] + [f"{lab} {x}" for lab, _, _ in arms for x in ("mean", "s.d.")] + [f"{lab} vs {arms[0][0]}" for lab, _, _ in arms[1:]] \
        + ([f"{arms[2][0]} vs {arms[1][0]}"] if n > 2 else []) + [f"{lab} wins vs {arms[0][0]}" for lab, _, _ in arms[1:]] + ["Notes"]
    ws.cell(r, 1, f"{title}; seeds {', '.join(map(str, seeds))}; anomaly rule (any junction > {thr:.0f} s in any arm): " + (f"seeds {bad}" if bad else "no hits")).font = TITLE
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol); ws.cell(r, 1).fill = TFILL; r += 1
    for c, h in enumerate(hdr, 1): ws.cell(r, c, h)
    style_row(ws, r, ncol=ncol, bold=True, fill=HFILL); r += 1
    key = {}
    def block(m, label, jlabel, tag):
        nonlocal r
        vals = {}
        for c in ("car", "bus", "ambulance", "all"):
            vals[c] = [np.array([X[s]["collector"]["summary"].get(f"eval/system/{c}/{m}", np.nan) for s in seeds]) for X in A]
        J = [sum(SH18[c] * WT[c] * np.nan_to_num(vals[c][i]) for c in SH18) for i in range(n)]
        for c in ("car", "bus", "ambulance", "all", "J"):
            arrs = J if c == "J" else vals[c]; lab = jlabel if c == "J" else f"{label} - {c}"
            ws.cell(r, 1, lab)
            for i, a in enumerate(arrs):
                ws.cell(r, 2 + 2 * i, float(np.nanmean(a))).number_format = "0.00"; ws.cell(r, 3 + 2 * i, float(np.nanstd(a))).number_format = "0.00"
            col = 2 + 2 * n
            for i in range(1, n):
                ws.cell(r, col, f'=IF(B{r}=0,"",{get_column_letter(2 + 2 * i)}{r}/B{r}-1)').number_format = "+0%;-0%;0%"; col += 1
            if n > 2:
                ws.cell(r, col, f'=IF(D{r}=0,"",F{r}/D{r}-1)').number_format = "+0%;-0%;0%"; col += 1
            for i in range(1, n):
                ok = ~np.isnan(arrs[0]) & ~np.isnan(arrs[i])
                ws.cell(r, col, f"{int((arrs[i] < arrs[0])[ok].sum())}/{int(ok.sum())}" + ("" if c == "J" else f" (ties {int((arrs[i] == arrs[0])[ok].sum())})")); col += 1
            if c == "ambulance": ws.cell(r, ncol, "1 ambulance per evaluation (probe route); do not cite the mean alone")
            style_row(ws, r, ncol=ncol)
            means = [float(np.nanmean(a)) for a in arrs]; ws.cell(r, 2 + 2 * int(np.argmin(means))).fill = GREEN
            key[(tag, c)] = r; r += 1
        car, bus, amb = key[(tag, "car")], key[(tag, "bus")], key[(tag, "ambulance")]
        ws.cell(r, 1, "Ordering gate (stopped time): amb <= bus <= car" if tag == "t" else "Ordering gate (stop events): amb <= bus <= car (amb n/a: bus <= car only)")
        for i in range(n):
            col = get_column_letter(2 + 2 * i)
            if tag == "t": ws.cell(r, 2 + 2 * i, f'=IF(AND({col}{amb}<={col}{bus},{col}{bus}<={col}{car}),"pass",IF(AND({col}{amb}-{col}{bus}<=5,{col}{bus}-{col}{car}<=5),"pass (tol. 5 s)","fail"))')
            else: ws.cell(r, 2 + 2 * i, f'=IF({col}{bus}<={col}{car},"pass (amb n/a)","fail (amb n/a)")')
        style_row(ws, r, ncol=ncol); r += 1
    block("avg_stopped_time_per_visit", "Stopped time / visit (s)", "J(531) weighted cost (per visit)", "t")
    block("avg_stop_events_per_visit", "Stop events / visit", "J(531) weighted stop events (per visit)", "e")
    r += 1
    # type swap, all arms
    ts = {}
    for i, (lab, _, _) in enumerate(arms):
        amb_pv, bus_pv, amb_ev, bus_ev = [], [], [], []
        for s in seeds:
            va = [A[i][s]["collector"]["vehicles"][k] for k in amb_ids if k in A[i][s]["collector"]["vehicles"]]
            vb = [Ab[i][s]["collector"]["vehicles"][k] for k in bus_ids if k in Ab[i][s]["collector"]["vehicles"]]
            amb_pv.append(_pv(va, "stopped_sec")); bus_pv.append(_pv(vb, "stopped_sec")); amb_ev.append(_pv(va, "stop_events")); bus_ev.append(_pv(vb, "stop_events"))
        ts[lab] = {"pv": (amb_pv, bus_pv), "ev": (amb_ev, bus_ev)}
    r = write_typeswap_arms(ws, r, seeds, ts, f"{title}: same route & departure time, ambulance vs bus probe, {len(seeds)} seeds")
    return r, key, bad

def write_typeswap_arms(ws, r, seeds, ts, title):
    ws.cell(r, 1, title).font = TITLE; ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8); ws.cell(r, 1).fill = TFILL; r += 1
    hdr = ["Vehicle / metric"] + [f"seed {s}" for s in seeds] + ["Mean", "S.d."]
    for c, h in enumerate(hdr, 1): ws.cell(r, c, h)
    style_row(ws, r, ncol=len(hdr), bold=True, fill=HFILL); r += 1
    n = len(seeds); last = get_column_letter(1 + n); lines = []
    for arm in ts:
        for metric, mk in (("Stopped time / visit (s)", "pv"), ("Stop events / visit", "ev")):
            for veh, idx in (("ambulance (probe route)", 0), ("bus (same route & time)", 1)):
                ws.cell(r, 1, f"{arm} - {veh} - {metric}")
                for k, v in enumerate(ts[arm][mk][idx]): ws.cell(r, 2 + k, float(v)).number_format = "0.00"
                ws.cell(r, 2 + n, f"=AVERAGE(B{r}:{last}{r})").number_format = "0.00"; ws.cell(r, 3 + n, f"=STDEV(B{r}:{last}{r})").number_format = "0.00"
                style_row(ws, r, ncol=len(hdr)); lines.append((arm, mk, idx, r)); r += 1
    for arm in ts:
        for metric, mk in (("stopped time", "pv"), ("stop events", "ev")):
            ra = [x[3] for x in lines if x[0] == arm and x[1] == mk and x[2] == 0][0]; rb = [x[3] for x in lines if x[0] == arm and x[1] == mk and x[2] == 1][0]
            ws.cell(r, 1, f"{arm} - {metric}: evaluations with ambulance < bus / ties")
            ws.cell(r, 2, f'=SUMPRODUCT(--(B{ra}:{last}{ra}<B{rb}:{last}{rb}))&" / "&SUMPRODUCT(--(B{ra}:{last}{ra}=B{rb}:{last}{rb}))&" (of {n})"')
            style_row(ws, r, ncol=2); r += 1
    return r + 1

def tail40_11h():
    """Dublin 11h tail-40 evaluations (ep605-800 every 5, eval seed 123): constants from plot_main_table_en_2026-09-22.py (S/EV)."""
    S = {"car": ((4.95, 1.24), (2.44, 0.32)), "bus": ((1.26, 0.18), (1.06, 0.14)), "amb": ((0.90, 0.69), (0.92, 0.48)), "all": ((4.75, 1.17), (2.38, 0.31)), "J": ((4.89, 1.18), (2.48, 0.31))}
    EV = {"car": ((0.57, 0.09), (0.27, 0.03)), "bus": ((0.18, 0.02), (0.15, 0.02)), "amb": ((0.14, 0.1), (0.15, 0.07)), "all": ((0.55, 0.09), (0.27, 0.03)), "J": ((0.566, 0.091), (0.281, 0.027))}
    L = {"car": "Stopped time / visit (s) - car", "bus": "Stopped time / visit (s) - bus", "amb": "Stopped time / visit (s) - ambulance", "all": "Stopped time / visit (s) - all", "J": "J(531) weighted cost (per visit)"}
    LE = {"car": "Stop events / visit - car", "bus": "Stop events / visit - bus", "amb": "Stop events / visit - ambulance", "all": "Stop events / visit - all", "J": "J(531) weighted stop events (per visit)"}
    rows_t = [[L[k], S[k][0][0], S[k][0][1], S[k][1][0], S[k][1][1], "", "40/40" if k == "J" else ""] for k in ("car", "bus", "amb", "all", "J")]
    rows_e = [[LE[k], EV[k][0][0], EV[k][0][1], EV[k][1][0], EV[k][1][1], "", "40/40" if k == "J" else ""] for k in ("car", "bus", "amb", "all", "J")]
    return rows_t, rows_e

def main():
    wb = Workbook(); ws0 = wb.active; ws0.title = "README"
    lines = ["Main comparison tables - 8STD (standard-phase template + DQN) vs GS-ENUM (enumerated phase sets + FRAP) - 2026-09-30", "",
             "Sheets: Overview | 1x1 | 1x3 | Dublin 02h | Dublin 11h | Dublin 18h. Each block: 5 stopped-time rows + ordering gate + 5 stop-event rows + ordering gate.",
             "Means and s.d. are measured data (copied from the json behind each PNG); relative differences, ordering gates, type-swap statistics and per-seed recomputations are formulas. Green cell = the better (lower) arm in that row.",
             "Protocol A (best checkpoint x evaluation seeds): each arm's best.pth (best sliding mean of 3 evaluations during training) is re-evaluated once per seed (123-132); training-collector metrics, 5 s sampling, per visit; 'GS wins' = number of seeds where GS is better (paired by seed).",
             "Protocol B (tail-40 evaluations): the 40 evaluations at episodes 605-800 (every 5) during training, evaluation seed 123 fixed - 40 different policies on one seed. Only the Dublin 11h tail-40 block uses this protocol (same data as the 11h block of main_comparison_8std_vs_frap_2026-09-22_en.png).",
             "A and B are not directly comparable: A measures the robustness of one fixed policy to traffic randomness; B measures late-training policy fluctuation.",
             "Anomaly exclusion only by symmetric, stated rules (any junction with all-class stopped time above the threshold in either arm -> the seed is removed from both arms). Each block title states whether the rule hit.",
             "J(531) = per-class per-visit metric weighted 5/3/1 (ambulance/bus/car) by each scenario's traffic shares (see plotting scripts).",
             "Ambulance sample: Dublin 02h has 2 ambulances per evaluation, 11h has 4, 18h has 1 (probe route); grids 1x1/1x3 have many more. Do not cite the Dublin ambulance means alone.",
             "Dublin 11h: the primary block is the USER-SELECTED protocol - seeds 128 and 132 removed from both arms, because in those two seeds the 8STD ambulance amb_0 was caught for 45 s / 90 s at junction 389281 by the Aungier Street right-turn queue spilling back from the unsignalised yield junction (an exogenous bottleneck outside the control set). The full 10-seed result is kept as a reference block; both appear in the Overview.",
             "Dublin 18h: PRIMARY block has three arms - 8STD (exp298 ep200), GS (exp300 ep145: split-slot net only) and GS-2 (exp301 ep95: exp300 + since-green observation + log-scale waiting + pressure-anchored phase score + TD-target floor -50); computed from the raw per-seed collector json with the same formulas as the PNG tables. Reference blocks below it: two-arm rerouted v2 (177 cars/h leave at the end of the Aungier approach instead of turning into Peter's Row; GS uses the split-slot enumerated net), best checkpoints exp298 ep200 (8STD, training stopped at ep253) vs exp300 ep145 (GS, training still running) - a snapshot, not a final result. A USER-SELECTED 5-seed block (seeds with zero ambulance stops in both arms) and the original-scenario blocks (exp289/290, 7 of 10 seeds anomalous by the >100 s rule) are kept as reference.",
             "Source files and generator: experiments/tools/figures/build_main_comparison_xlsx_2026-09-30.py; json paths are on line 2 of each sheet."]
    for i, t in enumerate(lines, 1): ws0.cell(i, 1, t).font = TITLE if i == 1 else N
    ws0.column_dimensions["A"].width = 180
    sheets = []
    ws, r, k = add_json(wb, "1x1", f"{FIG}/main_comparison_1x1_best10seeds_2026-09-30_en.json", amb_na=False, protocol="Protocol A"); sheets.append(("1x1", ws, k))
    ws, r, k = add_json(wb, "1x3", f"{FIG}/main_comparison_1x3_best10seeds_2026-09-30_en.json", amb_na=False, protocol="Protocol A"); sheets.append(("1x3", ws, k))
    ws, r, k = add_json(wb, "Dublin 02h", f"{FIG}/main_comparison_02h_best10seeds_2026-09-30_en.json", amb_na=True, protocol="Protocol A; ambulance route = 11h amb_1 route (2 ambulances)"); sheets.append(("Dublin 02h", ws, k))
    # Dublin 11h: user-selected (excl. 128/132) primary, full 10 seeds reference, tail-40 protocol B
    ws, r, k = add_json(wb, "Dublin 11h", f"{FIG}/main_comparison_11h_best10seeds_SELECTED_excl128_132_2026-09-30_en.json", amb_na=True,
                        protocol="Protocol A, USER-SELECTED: seeds 128 and 132 removed from both arms (8STD amb_0 caught at 389281 by Aungier St spillback, 45 s / 90 s)",
                        title_override="Dublin 11:00, best checkpoints x 8 seeds, seeds 128 & 132 removed (8STD: exp208 ep240, GS-ENUM: exp211 ep215)  [USER-SELECTED]",
                        ts_title="Dublin 11:00: same route & departure time, ambulance vs bus probe, best checkpoints, 8 seeds [USER-SELECTED]")
    ws, r, k2 = add_json(wb, "Dublin 11h", f"{FIG}/main_comparison_11h_best10seeds_2026-09-30_en.json", amb_na=True, protocol="Protocol A, all 10 seeds (reference)", ws=ws, r=r,
                         title_override="Dublin 11:00, best checkpoints x 10 seeds, all seeds (8STD: exp208 ep240, GS-ENUM: exp211 ep215)  [reference]", typeswap=False)
    rows_t, rows_e = tail40_11h()
    r, k3 = write_block(ws, r, "Dublin 11:00, tail-40 evaluations (episodes 605-800 every 5, evaluation seed 123)  (8STD: exp208, GS-ENUM: exp211)  [Protocol B]", rows_t, rows_e, amb_na=True,
                        note="Constants from plot_main_table_en_2026-09-22.py (S/EV) = the 11h block of main_comparison_8std_vs_frap_2026-09-22_en.png; 'GS wins 40/40' = J better in all 40 evaluations.")
    sheets.append(("Dublin 11h", ws, k + k2 + [("Dublin 11:00 tail-40 [Protocol B]", k3)]))
    # Dublin 18h: THREE-ARM primary block (8STD exp298 | GS exp300 | GS-2 exp301) computed from raw collector json, then 2-arm reference blocks
    ws = wb.create_sheet("Dublin 18h"); r = 1
    ws.cell(r, 1, "Dublin 18h: rerouted scenario v2 (177 cars/h leave at the Aungier approach end); GS arms on the split-slot enumerated net; SNAPSHOT (exp300 / exp301 still training)").font = TITLE; r += 1
    ws.cell(r, 1, "GS = exp300 (split-slot net, no other change); GS-2 = exp301 = exp300 + since-green observation + log-scale waiting + pressure-anchored phase score + TD-target floor -50. Best checkpoints frozen 2026-09-30 evening: 8STD exp298 ep200, GS exp300 ep145, GS-2 exp301 ep95.").font = NOTE; r += 2
    ARMS = [("8STD", 298, 200), ("GS (exp300)", 300, 145), ("GS-2 (exp301)", 301, 95)]
    r, k3a, bad = three_arm_block(ws, r, "Dublin 18:00 rerouted v2, best checkpoints x 10 seeds, three arms  [PRIMARY, snapshot]", ARMS)
    three_arm_keys = [("Dublin 18:00 rerouted v2, three arms, 10 seeds [PRIMARY]", k3a)]
    if bad:
        r, k3b, _ = three_arm_block(ws, r, f"Dublin 18:00 rerouted v2, three arms, excl. anomalous seeds {bad}", ARMS, seeds_keep=[x for x in range(123, 133) if x not in bad])
        three_arm_keys.append((f"Dublin 18:00 rerouted v2, three arms, excl. anomalous seeds {bad}", k3b))
    r += 1
    ws, r, k = add_json(wb, "Dublin 18h", f"{FIG}/main_comparison_18h_best10seeds_reroute_v2_2026-09-30_en.json", amb_na=True, ws=ws, r=r,
                        protocol="Protocol A; rerouted scenario v2 (Aungier right-turn cars leave at the approach end), GS on the split-slot enumerated net; SNAPSHOT: exp300 still training",
                        title_override="Dublin 18:00, rerouted scenario v2, best checkpoints x 10 seeds, two arms (8STD: exp298 ep200, GS-ENUM: exp300 ep145); anomaly rule (>100 s): no hits  [reference: same data as the PNG]",
                        ts_title="Dublin 18:00 rerouted v2: same route & departure time, ambulance vs bus probe, best checkpoints, 10 seeds")
    ws, r, k2 = add_json(wb, "Dublin 18h", f"{FIG}/main_comparison_18h_best10seeds_reroute_v2_SELECTED_ambzero5_2026-09-30_en.json", amb_na=True, ws=ws, r=r,
                         protocol="Protocol A, USER-SELECTED: 5 seeds with zero ambulance stops in both arms", typeswap=False,
                         title_override="Dublin 18:00, rerouted scenario v2, 5 seeds with zero ambulance stops in both arms (123, 126, 128, 131, 132)  [USER-SELECTED]")
    ws, r, k3 = add_json(wb, "Dublin 18h", f"{FIG}/main_comparison_18h_best10seeds_2026-09-30_en.json", amb_na=True, ws=ws, r=r,
                         protocol="Protocol A, ORIGINAL 18h scenario (exogenous Aungier/Peter's Row yield lock present), best checkpoints exp289 ep270 / exp290 ep75 (reference)",
                         title_override=["Dublin 18:00, ORIGINAL scenario, best checkpoints x 10 seeds (8STD: exp289 ep270, GS-ENUM: exp290 ep75)  [reference]",
                                         "Dublin 18:00, ORIGINAL scenario, excl. 7 anomalous seeds (any junction > 100 s in either arm: 124, 125, 127, 128, 130, 131, 132), 3 seeds  [reference]"],
                         ts_title="Dublin 18:00 ORIGINAL scenario: same route & departure time, ambulance vs bus probe, best checkpoints, 10 seeds [reference]")
    sheets.append(("Dublin 18h", ws, k + k2 + k3)); THREE = (ws, three_arm_keys)
    # Overview
    ov = wb.create_sheet("Overview", 1); r = 1
    ov.cell(r, 1, "Overview: all-class stopped time / visit and J(531) per block (formulas referencing the scenario sheets)").font = TITLE; r += 1
    hdr = ["Scenario / block", "Protocol", "8STD all-class stopped time / visit (s)", "GS-ENUM all-class stopped time / visit (s)", "Rel. diff", "8STD J(531)", "GS-ENUM J(531)", "J rel. diff", "GS wins (J)", "Ordering gate 8STD (stopped time)", "Ordering gate GS (stopped time)"]
    for c, h in enumerate(hdr, 1): ov.cell(r, c, h)
    style_row(ov, r, ncol=len(hdr), bold=True, fill=HFILL); r += 1
    for name, ws, keys in sheets:
        for t, key in keys:
            ra, rj = key[("t", "all")], key[("t", "J")]; q = f"'{ws.title}'!"; gate_row = rj + 1
            proto = "B" if "tail-40" in t else "A"
            vals = [f"{name}: {t}", proto, f"={q}B{ra}", f"={q}D{ra}", f'=IF(C{r}=0,"",D{r}/C{r}-1)', f"={q}B{rj}", f"={q}D{rj}", f'=IF(F{r}=0,"",G{r}/F{r}-1)', f"={q}G{rj}", f"={q}B{gate_row}", f"={q}D{gate_row}"]
            for c, v in enumerate(vals, 1): ov.cell(r, c, v)
            for c in (3, 4, 6, 7): ov.cell(r, c).number_format = "0.00"
            for c in (5, 8): ov.cell(r, c).number_format = "+0%;-0%;0%"
            style_row(ov, r, ncol=len(hdr)); r += 1
    r += 1; ws3, keys3 = THREE
    ov.cell(r, 1, "Dublin 18h three-arm block (8STD | GS exp300 | GS-2 exp301): all-class stopped time / visit and J(531)").font = TITLE; r += 1
    hdr3 = ["Block", "Protocol", "8STD all", "GS (exp300) all", "GS-2 (exp301) all", "GS vs 8STD", "GS-2 vs 8STD", "8STD J", "GS J", "GS-2 J", "GS wins (J)", "GS-2 wins (J)"]
    for c, h in enumerate(hdr3, 1): ov.cell(r, c, h)
    style_row(ov, r, ncol=len(hdr3), bold=True, fill=HFILL); r += 1
    for t, key in keys3:
        ra, rj = key[("t", "all")], key[("t", "J")]; q = f"'{ws3.title}'!"
        vals = [t, "A", f"={q}B{ra}", f"={q}D{ra}", f"={q}F{ra}", f'=IF(C{r}=0,"",D{r}/C{r}-1)', f'=IF(C{r}=0,"",E{r}/C{r}-1)', f"={q}B{rj}", f"={q}D{rj}", f"={q}F{rj}", f"={q}K{rj}", f"={q}L{rj}"]
        for c, v in enumerate(vals, 1): ov.cell(r, c, v)
        for c in (3, 4, 5, 8, 9, 10): ov.cell(r, c).number_format = "0.00"
        for c in (6, 7): ov.cell(r, c).number_format = "+0%;-0%;0%"
        style_row(ov, r, ncol=len(hdr3)); r += 1
    ov.column_dimensions["A"].width = 110; ov.column_dimensions["B"].width = 9
    for c in "CDEFGHIJK": ov.column_dimensions[c].width = 20
    ov.freeze_panes = "C3"
    os.makedirs(os.path.dirname(OUT), exist_ok=True); wb.save(OUT); print("saved", OUT)

if __name__ == "__main__": main()
